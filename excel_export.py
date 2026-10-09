"""Build a private Excel snapshot entirely in memory for Telegram delivery."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
import math
import re
from typing import TYPE_CHECKING
import warnings

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps, UnidentifiedImageError

if TYPE_CHECKING:
    from reporting import ReportSnapshot


MOSCOW = timezone(timedelta(hours=3))
HEADER_ROW = 5
DATE_FORMAT = "dd.mm.yyyy hh:mm:ss"
DURATION_FORMAT = "[h]:mm:ss"
GUEST_HEADERS = (
    "№ базы", "База", "ID гостя", "Фото", "Имя", "Телефон", "Статус входа",
    "Причина", "Комментарий", "Всего времени", "Завершённых посещений",
    "Время завершённых посещений", "Текущий визит", "Начало текущего визита (МСК)",
    "Время текущего визита", "Создано (МСК)", "Изменено (МСК)",
    "Создал (Telegram ID)", "Изменил (Telegram ID)", "Версия карточки",
    "Нормализованный телефон", "Фото (Telegram file_id)",
)
VISIT_HEADERS = (
    "№ базы", "База", "ID гостя", "Имя", "Телефон", "ID посещения",
    "Начало (МСК)", "Окончание (МСК)", "Длительность", "Состояние",
    "Начал (Telegram ID)", "Остановил (Telegram ID)",
)

_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_EMERALD = "123F35"
_GOLD = "E4C16E"
_PALE = "F0F6F3"
_LINE = "D9E5DE"


def _text(value: str) -> str:
    # Replace characters XML cannot represent; preserve all other user text,
    # including newlines, emoji and strings beginning with formula characters.
    # XML/Excel normalize line endings to LF; do so explicitly for consistent
    # output whether openpyxl uses its standard writer or optional lxml writer.
    return _ILLEGAL_XML.sub("\ufffd", value).replace("\r\n", "\n").replace("\r", "\n")


def _write(cell: Cell, value: object, *, identifier: bool = False) -> None:
    if isinstance(value, str):
        cleaned = _text(value)
        # openpyxl silently slices strings to 32,767 Python characters. Excel
        # counts UTF-16 units, so reject oversized legacy values beforehand;
        # a failed export is preferable to a file with silently lost card data.
        if len(cleaned.encode("utf-16-le")) // 2 > 32767:
            raise ValueError("Текст карточки превышает предел ячейки Excel (32767 символов).")
        cell.value = cleaned
        # Assignment alone treats '=' as a formula and '#N/A' as an error.
        # Explicit string cells preserve the original input without evaluation.
        cell.data_type = "s"
        if identifier:
            cell.number_format = "@"
    else:
        cell.value = value
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _local_time(value: str | datetime | None) -> datetime | None:
    if value is None:
        return None
    parsed = datetime.fromisoformat(value) if isinstance(value, str) else value
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise ValueError("Время выгрузки должно содержать часовой пояс.")
    # Excel stores timezone-free dates. Every date heading explicitly states
    # Moscow time, matching the guest cards and visit history in the bot.
    return parsed.astimezone(MOSCOW).replace(tzinfo=None)


def prepare_photo_thumbnail(photo: bytes) -> bytes:
    """Shrink a stored photo before retaining it in a multi-database snapshot.

    Invalid or unsafe image data returns empty bytes, which the workbook marks
    explicitly. This helper and the workbook builder never write image files.
    """
    if not photo:
        return b""
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(BytesIO(photo)) as original:
                # A Telegram photo is normally much smaller. Refuse excessive
                # decoded dimensions before allocating hundreds of MB on hosting.
                if original.width * original.height > 20_000_000:
                    return b""
                original.thumbnail((112, 112), Image.Resampling.LANCZOS)
                corrected = ImageOps.exif_transpose(original)
                rgba = corrected.convert("RGBA")
                rgb = Image.new("RGB", rgba.size, "white")
                rgb.paste(rgba, mask=rgba.getchannel("A"))
                with BytesIO() as buffer:
                    rgb.save(buffer, format="JPEG", quality=80, optimize=True)
                    return buffer.getvalue()
    except (UnidentifiedImageError, OSError, ValueError,
            Image.DecompressionBombError, Image.DecompressionBombWarning):
        # Legacy databases may contain invalid photo bytes. Keep the guest and
        # visibly mark the missing thumbnail instead of dropping the record.
        return b""


def _thumbnail(photo: bytes) -> ExcelImage | None:
    encoded = prepare_photo_thumbnail(photo)
    return ExcelImage(BytesIO(encoded)) if encoded else None


def _prepare_sheet(sheet, headers: tuple[str, ...], widths: tuple[float, ...],
                   snapshot: ReportSnapshot, *, title: str) -> None:
    last = get_column_letter(len(headers))
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.tabColor = _EMERALD
    sheet.merge_cells(f"A1:{last}1")
    _write(sheet["A1"], title)
    sheet["A1"].font = Font(name="Calibri", size=18, bold=True, color=_EMERALD)
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells("A2:D2")
    _write(sheet["A2"], "Выгрузка (МСК)")
    _write(sheet["E2"], _local_time(snapshot.generated_at))
    sheet["E2"].number_format = DATE_FORMAT
    sheet["A2"].font = Font(name="Calibri", bold=True, color=_EMERALD)
    sheet.merge_cells(f"A3:{last}3")
    names = "; ".join(f"{index}. {name}" for index, name in enumerate(snapshot.sources, 1))
    _write(sheet["A3"], f"Подключено баз: {len(snapshot.sources)}. {names}")
    sheet["A3"].font = Font(name="Calibri", size=10, color="52635B")
    sheet.row_dimensions[3].height = 30
    for column, (header, width) in enumerate(zip(headers, widths), 1):
        cell = sheet.cell(HEADER_ROW, column)
        _write(cell, header)
        cell.font = Font(name="Calibri", size=11, bold=True, color=_GOLD)
        cell.fill = PatternFill("solid", fgColor=_EMERALD)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[HEADER_ROW].height = 42
    sheet.freeze_panes = "E6"
    sheet.print_title_rows = "1:5"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def _style_row(sheet, row: int, values: tuple[object, ...], *, identifiers: set[int],
               dates: set[int], durations: set[int]) -> None:
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column)
        _write(cell, value, identifier=column in identifiers)
        cell.font = Font(name="Calibri", size=11, color="17251F")
        if row % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=_PALE)
        cell.border = Border(bottom=Side(style="hair", color=_LINE))
        if column in dates:
            cell.number_format = DATE_FORMAT
        elif column in durations:
            cell.number_format = DURATION_FORMAT
    sheet.row_dimensions[row].height = 32


def build_guest_workbook(snapshot: ReportSnapshot) -> bytes:
    """Export every supplied card, thumbnail and visit without paging or disk I/O.

    Durations include active visits up to ``snapshot.generated_at``. Duplicate
    guest IDs across sources remain separate through the source index/name.
    """
    workbook = Workbook()
    workbook.properties.creator = "PC Arena"
    workbook.properties.title = "Карточки гостей PC Arena"
    workbook.properties.description = "Гости и посещения. Все даты — московское время."
    workbook.properties.created = snapshot.generated_at.astimezone(timezone.utc).replace(tzinfo=None)
    workbook.properties.modified = workbook.properties.created
    guests = workbook.active
    guests.title = "Гости"
    visits = workbook.create_sheet("Посещения")
    _prepare_sheet(guests, GUEST_HEADERS,
                   (9, 22, 13, 21, 28, 24, 20, 52, 52, 20, 18, 23, 18, 24,
                    20, 24, 24, 24, 24, 15, 24, 38), snapshot,
                   title="PC Arena — карточки гостей")
    _prepare_sheet(visits, VISIT_HEADERS,
                   (9, 22, 13, 28, 24, 16, 24, 24, 20, 18, 24, 24), snapshot,
                   title="PC Arena — история посещений")

    visit_row = HEADER_ROW + 1
    for row, record in enumerate(snapshot.rows, HEADER_ROW + 1):
        guest, summary = record.guest, record.summary
        photo = _thumbnail(guest.photo)
        photo_label = "" if photo else ("Фото не удалось встроить" if guest.photo
                                       else "Фото отсутствует или повреждено")
        values = (
            record.source_index, record.source_name, guest.id, photo_label,
            guest.name, guest.phone,
            "Вход закрыт" if guest.entry_status == "closed" else "Вход открыт",
            guest.entry_reason, guest.comment, summary.total_seconds / 86400,
            summary.completed_count, summary.completed_seconds / 86400,
            "Идёт" if summary.active else "Нет",
            _local_time(summary.active.started_at) if summary.active else None,
            summary.current_seconds / 86400,
            _local_time(guest.created_at), _local_time(guest.updated_at),
            str(guest.created_by), str(guest.updated_by), guest.version,
            guest.phone_key, guest.photo_file_id,
        )
        _style_row(guests, row, values, identifiers={6, 18, 19, 21, 22},
                   dates={14, 16, 17}, durations={10, 12, 15})
        lines = max(sum(max(1, math.ceil(len(line) / 48)) for line in text.split("\n"))
                    for text in (guest.comment, guest.entry_reason))
        guests.row_dimensions[row].height = min(409, max(90, lines * 15 + 6))
        guests.cell(row, 7).font = Font(name="Calibri", size=11, bold=True,
                                       color="A02525" if guest.entry_status == "closed" else "156245")
        if photo:
            guests.add_image(photo, f"D{row}")

        for visit in record.visits:
            visit_values = (
                record.source_index, record.source_name, guest.id, guest.name,
                guest.phone, visit.id, _local_time(visit.started_at),
                _local_time(visit.stopped_at),
                visit.duration_seconds(snapshot.generated_at) / 86400,
                "Завершено" if visit.stopped_at else "Идёт",
                str(visit.started_by), str(visit.stopped_by) if visit.stopped_by is not None else None,
            )
            _style_row(visits, visit_row, visit_values, identifiers={5, 11, 12},
                       dates={7, 8}, durations={9})
            visit_row += 1

    guests.auto_filter.ref = f"A{HEADER_ROW}:V{max(HEADER_ROW, guests.max_row)}"
    visits.auto_filter.ref = f"A{HEADER_ROW}:L{max(HEADER_ROW, visits.max_row)}"
    with BytesIO() as output:
        try:
            workbook.save(output)
            return output.getvalue()
        finally:
            workbook.close()
