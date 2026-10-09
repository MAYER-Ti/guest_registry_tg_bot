from dataclasses import replace
from datetime import datetime, timedelta, timezone
from io import BytesIO
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook
from PIL import Image

from excel_export import (
    DATE_FORMAT, DURATION_FORMAT, GUEST_HEADERS, HEADER_ROW, VISIT_HEADERS,
    build_guest_workbook, prepare_photo_thumbnail,
)
from reporting import ReportGuest, ReportSnapshot
from storage import Guest, Visit, VisitSummary


NOW = datetime(2026, 10, 9, 12, 30, 40, tzinfo=timezone.utc)


def make_photo(color="green"):
    with BytesIO() as buffer:
        Image.new("RGB", (320, 240), color).save(buffer, format="PNG")
        return buffer.getvalue()


def make_guest(**fields):
    data = dict(id=1, name="Алёна Иванова", phone="+7 (999) 123-45-67",
                phone_key="79991234567", comment="Постоянная гостья\nЛюбит покер 😀",
                photo=make_photo(), photo_file_id="telegram-photo-id",
                created_by=7011639945, updated_by=987654321012345,
                created_at="2026-10-08T20:30:00+00:00",
                updated_at="2026-10-09T11:00:00+00:00", version=3,
                entry_status="open", entry_reason="")
    data.update(fields)
    return Guest(**data)


def make_record(guest=None, *, source_index=1, source_name="Основная база",
                visits=(), summary=None):
    return ReportGuest(source_index, source_name, guest or make_guest(),
                       summary or VisitSummary(None, 0, 0, 0), tuple(visits))


def snapshot(*rows, sources=("Основная база",)):
    return ReportSnapshot(NOW, tuple(sources), tuple(rows))


def fields(sheet, row=HEADER_ROW + 1):
    return {sheet.cell(HEADER_ROW, column).value: sheet.cell(row, column)
            for column in range(1, sheet.max_column + 1)}


class ExcelExportTests(unittest.TestCase):
    def export(self, data):
        raw = build_guest_workbook(data)
        self.assertIsInstance(raw, bytes)
        self.assertTrue(raw.startswith(b"PK"))
        workbook = load_workbook(BytesIO(raw))
        self.addCleanup(workbook.close)
        return raw, workbook

    def test_all_cards_sources_photos_and_full_visit_history(self):
        completed = Visit(1, 1, "2026-10-08T00:00:00+00:00",
                          "2026-10-09T03:00:00+00:00", 11, 22)
        active = Visit(2, 1, "2026-10-09T12:00:00+00:00", None, 22, None)
        rows = (
            make_record(visits=(completed, active),
                        summary=VisitSummary(active, 1, 27 * 3600, 30 * 60 + 40)),
            make_record(make_guest(name="Иван Петров", entry_status="closed",
                                   entry_reason="Нарушил правила", photo=make_photo("gold")),
                        source_index=2, source_name="Вторая база"),
        )
        raw, workbook = self.export(snapshot(*rows, sources=("Основная база", "Вторая база")))
        self.assertEqual(workbook.sheetnames, ["Гости", "Посещения"])
        guests = workbook["Гости"]
        self.assertEqual(guests.max_row, HEADER_ROW + 2)
        first, second = fields(guests), fields(guests, HEADER_ROW + 2)
        for key, expected in {
            "№ базы": 1, "База": "Основная база", "ID гостя": 1,
            "Имя": rows[0].guest.name, "Телефон": rows[0].guest.phone,
            "Статус входа": "Вход открыт", "Комментарий": rows[0].guest.comment,
            "Нормализованный телефон": rows[0].guest.phone_key,
            "Фото (Telegram file_id)": rows[0].guest.photo_file_id,
            "Версия карточки": 3,
        }.items():
            self.assertEqual(first[key].value, expected)
        self.assertEqual(second["№ базы"].value, 2)
        self.assertEqual(second["ID гостя"].value, 1)
        self.assertEqual(second["Статус входа"].value, "Вход закрыт")
        self.assertEqual(second["Причина"].value, "Нарушил правила")
        self.assertEqual(second["Имя"].value, "Иван Петров")
        self.assertEqual(guests.auto_filter.ref, "A5:V7")
        self.assertEqual(guests.freeze_panes, "E6")
        self.assertEqual(len(guests._images), 2)
        self.assertEqual(workbook["Посещения"].max_row, HEADER_ROW + 2)
        history = fields(workbook["Посещения"])
        self.assertEqual(history["ID посещения"].value, 1)
        self.assertEqual(history["Длительность"].value, timedelta(hours=27))
        self.assertEqual(history["Состояние"].value, "Завершено")
        self.assertEqual(history["Начал (Telegram ID)"].value, "11")
        self.assertEqual(history["Остановил (Telegram ID)"].value, "22")
        last = fields(workbook["Посещения"], HEADER_ROW + 2)
        self.assertIsNone(last["Окончание (МСК)"].value)
        self.assertEqual(last["Длительность"].value, timedelta(minutes=30, seconds=40))
        self.assertEqual(last["Состояние"].value, "Идёт")
        with ZipFile(BytesIO(raw)) as archive:
            media = [name for name in archive.namelist() if name.startswith("xl/media/")]
            self.assertEqual(len(media), 2)
            for name in media:
                with Image.open(BytesIO(archive.read(name))) as image:
                    self.assertLessEqual(max(image.size), 112)
                    self.assertEqual(image.format, "JPEG")

    def test_empty_workbook_keeps_headers_filters_and_empty_source_names(self):
        _, workbook = self.export(snapshot(sources=("Основная база", "Пустая база")))
        for title, headers, last_column in (("Гости", GUEST_HEADERS, "V"),
                                            ("Посещения", VISIT_HEADERS, "L")):
            sheet = workbook[title]
            self.assertEqual(sheet.max_row, HEADER_ROW)
            self.assertEqual(tuple(cell.value for cell in sheet[HEADER_ROW]), headers)
            self.assertEqual(sheet.auto_filter.ref, f"A5:{last_column}5")
            self.assertIn("2. Пустая база", sheet["A3"].value)
            self.assertEqual(len(sheet._images), 0)

    def test_formula_and_error_inputs_are_literal_text_without_external_links(self):
        guest = make_guest(name='=HYPERLINK("https://example.org","click")',
                           phone="+79990000000", comment="=1+1",
                           entry_reason="@SUM(A1:A2)", photo_file_id="#N/A")
        row = make_record(guest, source_name="=1+1")
        raw, workbook = self.export(snapshot(row, sources=("=1+1",)))
        exported = fields(workbook["Гости"])
        for heading, original in (("База", row.source_name), ("Имя", guest.name),
                                  ("Телефон", guest.phone), ("Комментарий", guest.comment),
                                  ("Причина", guest.entry_reason),
                                  ("Фото (Telegram file_id)", guest.photo_file_id)):
            self.assertEqual(exported[heading].value, original)
            self.assertEqual(exported[heading].data_type, "s")
            self.assertIsNone(exported[heading].hyperlink)
        with ZipFile(BytesIO(raw)) as archive:
            self.assertFalse(any(name.startswith("xl/externalLinks/") for name in archive.namelist()))
            for name in archive.namelist():
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                    root = ElementTree.fromstring(archive.read(name))
                    self.assertFalse(root.findall(".//{*}f"))

    def test_illegal_xml_is_replaced_and_maximum_unicode_text_is_preserved(self):
        comment = "😀" * 2995 + "\n\t\rаб"
        reason = "Текст\x00\x08\x0b\ud800\ufffe\uffff\nПродолжение"
        guest = make_guest(comment=comment, entry_reason=reason)
        _, workbook = self.export(snapshot(make_record(guest)))
        exported = fields(workbook["Гости"])
        self.assertEqual(exported["Комментарий"].value, comment.replace("\r", "\n"))
        self.assertEqual(len(exported["Комментарий"].value), 3000)
        self.assertEqual(exported["Причина"].value, "Текст" + "\ufffd" * 6 + "\nПродолжение")

    def test_typed_moscow_dates_elapsed_time_over_24h_and_text_identifiers(self):
        active = Visit(9, 1, (NOW - timedelta(hours=25, seconds=1)).isoformat(), None, 11, None)
        summary = VisitSummary(active, 3, 50 * 3600, 25 * 3600 + 1)
        guest = make_guest(created_at="2026-10-08T23:30:00-04:00")
        _, workbook = self.export(snapshot(make_record(guest, visits=(active,), summary=summary)))
        sheet = workbook["Гости"]
        exported = fields(sheet)
        self.assertEqual(sheet["E2"].value, datetime(2026, 10, 9, 15, 30, 40))
        self.assertEqual(exported["Создано (МСК)"].value, datetime(2026, 10, 9, 6, 30))
        self.assertEqual(exported["Создано (МСК)"].number_format, DATE_FORMAT)
        self.assertEqual(exported["Всего времени"].value, timedelta(hours=75, seconds=1))
        self.assertEqual(exported["Всего времени"].number_format, DURATION_FORMAT)
        self.assertEqual(exported["Завершённых посещений"].value, 3)
        self.assertEqual(exported["Завершённых посещений"].data_type, "n")
        self.assertEqual(exported["Текущий визит"].value, "Идёт")
        for heading, expected in (("Телефон", guest.phone),
                                  ("Создал (Telegram ID)", str(guest.created_by)),
                                  ("Изменил (Telegram ID)", str(guest.updated_by))):
            self.assertEqual(exported[heading].value, expected)
            self.assertEqual(exported[heading].data_type, "s")
            self.assertEqual(exported[heading].number_format, "@")

    def test_oversized_legacy_text_fails_instead_of_silent_excel_truncation(self):
        for field in ("comment", "entry_reason"):
            for text in ("а" * 32768, "😀" * 16384):
                with self.subTest(field=field, utf16_units=len(text.encode("utf-16-le")) // 2):
                    guest = make_guest(**{field: text})
                    with self.assertRaisesRegex(ValueError, "32767"):
                        build_guest_workbook(snapshot(make_record(guest)))
        # The exact UTF-16 boundary remains intact, including supplementary
        # Unicode characters. Ordinary 3,000-character bot fields fit easily.
        text = "😀" * 16383 + "а"
        _, workbook = self.export(snapshot(make_record(make_guest(comment=text))))
        self.assertEqual(fields(workbook["Гости"])["Комментарий"].value, text)

    def test_broken_and_absent_photos_keep_cards_and_explicit_markers(self):
        _, workbook = self.export(snapshot(make_record(make_guest(photo=b"fake-photo")),
                                           make_record(make_guest(id=2, photo=b""))))
        sheet = workbook["Гости"]
        self.assertEqual(sheet.max_row, HEADER_ROW + 2)
        self.assertEqual(fields(sheet)["Фото"].value, "Фото не удалось встроить")
        self.assertEqual(fields(sheet, HEADER_ROW + 2)["Фото"].value, "Фото отсутствует или повреждено")
        self.assertEqual(len(sheet._images), 0)

    def test_preparing_photo_shrinks_before_snapshot_without_changing_card_text(self):
        photo = prepare_photo_thumbnail(make_photo())
        self.assertTrue(photo.startswith(b"\xff\xd8"))
        self.assertLess(len(photo), 4096)
        with Image.open(BytesIO(photo)) as image:
            self.assertLessEqual(max(image.size), 112)
        self.assertEqual(prepare_photo_thumbnail(b"broken"), b"")
        guest = make_guest(photo=photo)
        _, workbook = self.export(snapshot(make_record(guest)))
        self.assertEqual(fields(workbook["Гости"])["Имя"].value, guest.name)
        self.assertEqual(len(workbook["Гости"]._images), 1)

    def test_all_rows_and_visits_are_exported_without_bot_page_limit(self):
        base = make_guest(photo=b"")
        rows = []
        for index in range(1, 102):
            guest = replace(base, id=index, name=f"Гость {index}")
            history = tuple(Visit(index * 10 + n, index, "2026-10-09T10:00:00+00:00",
                                  "2026-10-09T11:00:00+00:00", 11, 22)
                            for n in range(7))
            rows.append(make_record(guest, visits=history, summary=VisitSummary(None, 7, 7 * 3600, 0)))
        _, workbook = self.export(snapshot(*rows))
        self.assertEqual(workbook["Гости"].max_row, HEADER_ROW + 101)
        self.assertEqual(workbook["Посещения"].max_row, HEADER_ROW + 707)
        self.assertEqual(fields(workbook["Гости"], HEADER_ROW + 101)["Имя"].value, "Гость 101")
        self.assertEqual(fields(workbook["Посещения"], HEADER_ROW + 707)["ID посещения"].value, 1016)


if __name__ == "__main__":
    unittest.main()
