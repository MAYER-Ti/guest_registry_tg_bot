"""Private reporting buttons exercised through Telegram updates and real SQLite."""
from __future__ import annotations

import base64
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
import shutil
import unittest
from unittest.mock import AsyncMock, call, patch
from uuid import uuid4
from zipfile import ZipFile

from aiogram import Bot
from aiogram.exceptions import TelegramRetryAfter
from aiogram.methods import AnswerCallbackQuery, SendDocument, SendMessage
from aiogram.types import CallbackQuery, Chat, Message, Update, User
from openpyxl import load_workbook
from PIL import Image

import main
from config import Config
from reporting import ReportError
from storage import Store
from test_bot import RecordingSession


# Valid image bytes: exports must embed the stored photo, without Telegram I/O.
PHOTO = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a63sAAAAASUVORK5CYII="
)


class DocumentRecordingSession(RecordingSession):
    def __init__(self):
        super().__init__()
        self.retry_on_message = None
        self.message_attempts = []

    async def make_request(self, bot, method, timeout=None):
        if isinstance(method, SendMessage):
            self.message_attempts.append(method)
            if len(self.message_attempts) == self.retry_on_message:
                raise TelegramRetryAfter(method=method, message="Too Many Requests", retry_after=3)
        if isinstance(method, SendDocument):
            self.calls.append(method)
            return Message(
                message_id=10000 + len(self.calls), date=datetime.now(timezone.utc),
                chat=Chat(id=int(method.chat_id), type="private"),
                from_user=User(id=bot.id, is_bot=True, first_name="Guest notebook"),
                caption=method.caption,
            )
        return await super().make_request(bot, method, timeout=timeout)


class ReportButtonTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.runtime = Path.cwd() / ".test-runtime" / uuid4().hex
        self.runtime.mkdir(parents=True, mode=0o777)
        self.store = Store(self.runtime / "primary.sqlite3")
        self.session = DocumentRecordingSession()
        self.bot = Bot("123456789:ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijk", session=self.session)
        self.update_id = 0
        self.dp = None
        self.sleep_patch = patch.object(main.asyncio, "sleep", new_callable=AsyncMock)
        self.sleep_mock = self.sleep_patch.start()
        await self.configure()

    async def asyncTearDown(self):
        self.sleep_patch.stop()
        if self.dp:
            await self.dp.storage.close()
            await self.dp.fsm.events_isolation.close()
        await self.bot.session.close()
        shutil.rmtree(self.runtime)
        try:
            self.runtime.parent.rmdir()
        except OSError:
            pass

    async def configure(self, extra_paths=()):
        if self.dp:
            await self.dp.storage.close()
            await self.dp.fsm.events_isolation.close()
        self.dp = main.build_dispatcher(
            Config(token=self.bot.token, allowed_user_ids=frozenset({101, 102, 103, 104}),
                   db_path=self.store.path, extra_db_paths=tuple(extra_paths)), self.store,
        )

    async def message(self, text, *, user_id=101, group=False):
        self.update_id += 1
        incoming = Message(
            message_id=self.update_id, date=datetime.now(timezone.utc),
            chat=Chat(id=-999 if group else user_id, type="group" if group else "private"),
            from_user=User(id=user_id, is_bot=False, first_name="Staff"), text=text,
        )
        before = len(self.session.calls)
        await self.dp.feed_update(self.bot, Update(update_id=self.update_id, message=incoming))
        return self.session.calls[before:]

    async def callback(self, data):
        self.update_id += 1
        incoming = CallbackQuery(
            id=str(self.update_id), from_user=User(id=101, is_bot=False, first_name="Staff"),
            chat_instance="test", data=data,
            message=Message(
                message_id=300, date=datetime.now(timezone.utc),
                chat=Chat(id=101, type="private"),
                from_user=User(id=self.bot.id, is_bot=True, first_name="Bot"), text="Card",
            ),
        )
        before = len(self.session.calls)
        await self.dp.feed_update(self.bot, Update(update_id=self.update_id, callback_query=incoming))
        return self.session.calls[before:]

    def seed(self, *, store=None, index=0, name=None, closed=False, reason="", photo=PHOTO):
        return (store or self.store).add_guest(
            name=name or f"Гость {index:03d}", phone=f"+7999123{index:04d}",
            comment=f"Комментарий гостя {index}", photo=photo,
            photo_file_id=f"telegram-photo-{index}", actor_id=101,
            entry_status="closed" if closed else "open", entry_reason=reason,
        )

    @staticmethod
    def texts(calls):
        return "".join(call.text for call in calls if isinstance(call, SendMessage))

    def assert_main_menu(self, calls):
        menus = [call.reply_markup for call in calls
                 if getattr(call, "reply_markup", None) and
                 getattr(call.reply_markup, "keyboard", None)]
        self.assertTrue(menus, "Report should restore the main menu")
        labels = {button.text for row in menus[-1].keyboard for button in row}
        self.assertTrue({main.ADD, main.FIND, main.TOTAL, main.BLACKLIST, main.EXPORT} <= labels)

    @staticmethod
    def workbook_values(payload):
        workbook = load_workbook(BytesIO(payload), data_only=False)
        values = [value for sheet in workbook for row in sheet.iter_rows(values_only=True)
                  for value in row if value is not None]
        return workbook, values

    async def test_new_buttons_are_in_main_menu_and_empty_total_is_zero(self):
        started = await self.message("/start")
        self.assert_main_menu(started)
        count = await self.message(main.TOTAL)
        self.assertIn("Всего гостей: 0", self.texts(count))
        self.assert_main_menu(count)

    async def test_count_includes_each_card_and_connected_source_without_id_deduplication(self):
        self.seed(index=1)
        self.seed(index=2, closed=True, reason="Причина")
        extra = Store(self.runtime / "another # source.sqlite3")
        collision = self.seed(store=extra, index=1, name="Карточка второй базы")
        self.assertEqual(1, collision.id)
        await self.configure((extra.path,))
        calls = await self.message(main.TOTAL, user_id=104)
        self.assertIn("Всего гостей: 3", self.texts(calls))
        self.assert_main_menu(calls)
        self.assertTrue(all(call.protect_content for call in calls if isinstance(call, SendMessage)))

    async def test_blacklist_is_full_excludes_open_and_splits_long_unicode_reason_safely(self):
        entries = [self.seed(index=i, closed=True, reason=f"Причина {i:03d}")
                   for i in range(main.PAGE_SIZE + 3)]
        long_reason = "Запрет: " + "😀" * 2500
        long = self.seed(index=50, name="Ещё один гость 🦉", closed=True, reason=long_reason)
        opened = self.seed(index=99, name="Открытый гость НЕ ПОКАЗЫВАТЬ")
        calls = await self.message(main.BLACKLIST)
        messages = [call for call in calls if isinstance(call, SendMessage)]
        full = self.texts(calls)
        for guest in [*entries, long]:
            self.assertEqual(1, full.count(guest.name))
            self.assertIn(guest.phone, full)
            self.assertIn(guest.entry_reason, full)
        self.assertNotIn(opened.name, full)
        self.assertNotIn(opened.phone, full)
        self.assertGreater(len(messages), 1)
        self.assertTrue(all(call.protect_content for call in messages))
        self.assertTrue(all(len(call.text.encode("utf-16-le")) // 2 <= 4000 for call in messages))
        self.assert_main_menu(calls)
        self.assertFalse(any(getattr(call.reply_markup, "inline_keyboard", None) for call in messages))

    async def test_blacklist_includes_distinct_sources_with_colliding_ids(self):
        first = self.seed(index=1, name="Иван основной базы", closed=True, reason="Причина основной")
        extra = Store(self.runtime / "extra.sqlite3")
        second = self.seed(store=extra, index=1, name="Иван дополнительной базы", closed=True,
                           reason="Причина дополнительной")
        self.assertEqual(first.id, second.id)
        self.seed(store=extra, index=2, name="Разрешённый другой базы")
        await self.configure((extra.path,))
        calls = await self.message(main.BLACKLIST, user_id=103)
        text = self.texts(calls)
        self.assertIn(first.name, text)
        self.assertIn(second.name, text)
        self.assertIn("Основная база", text)
        self.assertIn("База 2", text)
        self.assertNotIn("Разрешённый другой базы", text)

    async def test_blacklist_retries_rate_limited_middle_chunk_without_skipping_any_guests(self):
        entries = [self.seed(index=index, closed=True,
                             reason=f"Причина {index}: " + "😀" * 2500)
                   for index in range(3)]
        self.session.retry_on_message = 2
        calls = await self.message(main.BLACKLIST)
        sent = [method for method in calls if isinstance(method, SendMessage)]
        attempts = self.session.message_attempts
        self.assertGreaterEqual(len(sent), 3)
        self.assertEqual(len(sent) + 1, len(attempts))
        self.assertEqual(attempts[1].text, attempts[2].text,
                         "The rate-limited chunk must be retried unchanged")
        self.assertEqual(attempts[1].reply_markup, attempts[2].reply_markup)
        full = self.texts(calls)
        for guest in entries:
            self.assertEqual(1, full.count(guest.name))
            self.assertEqual(1, full.count(guest.phone))
            self.assertIn(guest.entry_reason, full)
        self.assertTrue(all(method.protect_content for method in sent))
        self.assertTrue(all(len(method.text.encode("utf-16-le")) // 2 <= 4000 for method in sent))
        self.sleep_mock.assert_any_await(3.1)
        self.assertEqual(len(sent) - 1, self.sleep_mock.await_args_list.count(call(1.05)))
        self.assert_main_menu(calls)

    async def test_empty_blacklist_responds_and_restores_menu(self):
        self.seed(index=1)
        calls = await self.message(main.BLACKLIST)
        self.assertTrue(self.texts(calls))
        self.assertNotIn("Гость 001", self.texts(calls))
        self.assert_main_menu(calls)
        self.assertFalse(any(isinstance(call, SendDocument) for call in calls))

    async def test_export_sends_complete_valid_workbook_with_connected_cards_photos_and_visits(self):
        first = self.seed(index=1, name="Основная карточка", closed=True, reason="Нарушил правила")
        started_at = datetime(2026, 10, 8, 9, 0, tzinfo=timezone.utc)
        visit = self.store.start_visit(first.id, first.version, 101, now=started_at)
        self.store.stop_visit(first.id, visit.id, 102, now=started_at + timedelta(seconds=90))
        extra = Store(self.runtime / "extra.sqlite3")
        second = self.seed(store=extra, index=1, name="Дополнительная карточка")
        self.assertEqual(first.id, second.id)
        await self.configure((extra.path,))
        before_files = set(self.runtime.rglob("*"))
        calls = await self.message(main.EXPORT, user_id=102)
        documents = [call for call in calls if isinstance(call, SendDocument)]
        self.assertEqual(1, len(documents))
        document = documents[0]
        self.assertTrue(document.protect_content)
        self.assertTrue(document.document.filename.endswith(".xlsx"))
        self.assertIn("2", document.caption)
        self.assertIn("2", self.texts(calls) + document.caption)
        payload = document.document.data
        self.assertTrue(payload.startswith(b"PK"))
        workbook, values = self.workbook_values(payload)
        for expected in (first.name, second.name, first.phone, first.entry_reason,
                         first.comment, second.comment):
            self.assertIn(expected, values)
        self.assertIn(timedelta(seconds=90), values)
        self.assertIn("Основная база", values)
        self.assertIn("База 2", values)
        with ZipFile(BytesIO(payload)) as package:
            media = [name for name in package.namelist() if name.startswith("xl/media/")]
            self.assertEqual(2, len(media))
            for name in media:
                with Image.open(BytesIO(package.read(name))) as image:
                    image.load()
                    self.assertEqual((1, 1), image.size)
        self.assertEqual(2, sum(len(sheet._images) for sheet in workbook))
        self.assert_main_menu(calls)
        after_files = set(self.runtime.rglob("*"))
        self.assertFalse(any(path.suffix == ".xlsx" for path in after_files - before_files))

    async def test_empty_export_is_valid_workbook_and_protected_document(self):
        calls = await self.message(main.EXPORT)
        document = next(call for call in calls if isinstance(call, SendDocument))
        self.assertTrue(document.protect_content)
        workbook, values = self.workbook_values(document.document.data)
        self.assertEqual({"Гости", "Посещения"}, set(workbook.sheetnames))
        self.assertIn("Карточек: 0", document.caption)
        self.assertFalse(any(isinstance(value, int) for value in values))
        self.assert_main_menu(calls)

    async def test_report_buttons_cancel_in_progress_draft_and_stale_delete_confirmation(self):
        guest = self.seed(index=1)
        for action in (main.TOTAL, main.BLACKLIST, main.EXPORT):
            with self.subTest(action=action):
                await self.message(main.ADD)
                await self.message(action)
                state = await self.dp.fsm.get_context(bot=self.bot, chat_id=101, user_id=101).get_state()
                self.assertIsNone(state)
                self.assertEqual(1, self.store.count_search("Гость"))
        asked = await self.callback(f"delete:{guest.id}:{guest.version}")
        confirmation = next(button.callback_data for call in asked
                            for row in getattr(getattr(call, "reply_markup", None), "inline_keyboard", [])
                            for button in row if (button.callback_data or "").startswith("remove:"))
        await self.message(main.TOTAL)
        stale = await self.callback(confirmation)
        self.assertTrue(any(isinstance(call, AnswerCallbackQuery) and call.show_alert for call in stale))
        self.assertIsNotNone(self.store.get_guest(guest.id))

    async def test_report_buttons_reject_strangers_and_groups_before_reading_any_source(self):
        guest = self.seed(index=1)
        with patch.object(main, "report_snapshot", side_effect=AssertionError("Must not read")) as read:
            for action in (main.TOTAL, main.BLACKLIST, main.EXPORT):
                for user_id, group in ((999, False), (101, True)):
                    with self.subTest(action=action, user_id=user_id, group=group):
                        calls = await self.message(action, user_id=user_id, group=group)
                        self.assertFalse(any(isinstance(call, SendDocument) for call in calls))
                        self.assertNotIn(guest.name, self.texts(calls))
            read.assert_not_called()

    async def test_missing_connected_source_never_sends_partial_report_or_creates_database(self):
        secret_name = "Секретная карточка не должна попасть в отчёт"
        self.seed(index=1, name=secret_name, closed=True, reason="Секретная причина")
        missing = self.runtime / "private-token-folder" / "missing-source.sqlite3"
        await self.configure((missing,))
        for action in (main.TOTAL, main.BLACKLIST, main.EXPORT):
            with self.subTest(action=action):
                calls = await self.message(action)
                text = self.texts(calls)
                self.assertTrue(text)
                self.assertNotIn(secret_name, text)
                self.assertNotIn("Секретная причина", text)
                self.assertNotIn(str(missing), text)
                self.assertNotIn(missing.name, text)
                self.assertNotIn(self.bot.token, text)
                self.assertNotIn("Всего гостей: 1", text)
                self.assertFalse(any(isinstance(call, SendDocument) for call in calls))
                self.assert_main_menu(calls)
        self.assertFalse(missing.exists())

    async def test_raw_reporting_exception_is_not_echoed_to_chat_or_used_for_partial_data(self):
        secret = f"database-path {self.runtime} token {self.bot.token} guest-sensitive-data"
        with patch.object(main, "report_snapshot", side_effect=ReportError(secret)):
            for action in (main.TOTAL, main.BLACKLIST, main.EXPORT):
                with self.subTest(action=action):
                    calls = await self.message(action)
                    text = self.texts(calls)
                    self.assertTrue(text)
                    for sensitive in (self.bot.token, str(self.runtime), "guest-sensitive-data"):
                        self.assertNotIn(sensitive, text)
                    self.assertFalse(any(isinstance(call, SendDocument) for call in calls))
                    self.assert_main_menu(calls)

    async def test_oversized_workbook_sends_clear_error_without_document(self):
        self.seed(index=1)
        with patch.object(main, "MAX_EXPORT_BYTES", 32), \
                patch.object(main, "build_guest_workbook", return_value=b"x" * 33):
            calls = await self.message(main.EXPORT)
        self.assertFalse(any(isinstance(call, SendDocument) for call in calls))
        self.assertTrue(self.texts(calls))
        self.assert_main_menu(calls)


if __name__ == "__main__":
    unittest.main()
